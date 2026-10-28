"""docx_cover.py — 对 pandoc 生成的 docx 做后处理：封面、三线表、页眉、页码（供 build.py 调用）"""
from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK, BLUE, TEAL, GREY = RGBColor(0x08, 0x16, 0x50), RGBColor(0x1B, 0x3F, 0xB0), RGBColor(0x14, 0x9C, 0x8E), RGBColor(0x59, 0x60, 0x7A)
SIDE = Cm(2.4)


def _fonts(run, east="宋体", latin="Times New Roman"):
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), latin)
    rf.set(qn("w:eastAsia"), east)


def _run(p, text, size=10.5, bold=False, color=None, east="宋体", latin="Times New Roman", spacing=None):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    _fonts(r, east, latin)
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        r._r.get_or_add_rPr().append(sp)
    return r


def _cell_border(cell, **kw):
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcpr.append(b)
    for side, (val, sz, color) in kw.items():
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)


# ------------------------------------------------------------ 三线表
def style_tables(doc):
    """全部表格：三线表（上/下粗线，表头下细线）、表头浅蓝底加粗、字号 9.5、居中。"""
    for t in doc.tables:
        tblpr = t._tbl.tblPr
        for old in tblpr.findall(qn("w:tblBorders")):
            tblpr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for side, sz in (("top", 12), ("bottom", 12)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "081650")
            borders.append(e)
        for side in ("insideH",):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "2"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "D5DBEE")
            borders.append(e)
        tblpr.append(borders)
        jc = tblpr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            tblpr.append(jc)
        jc.set(qn("w:val"), "center")
        # 单元格内边距
        mar = OxmlElement("w:tblCellMar")
        for side, w in (("top", 40), ("bottom", 40), ("left", 80), ("right", 80)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(w)); e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblpr.append(mar)
        for ri, row in enumerate(t.rows):
            for cell in row.cells:
                if ri == 0:
                    tcpr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EAF0FF")
                    tcpr.append(shd)
                    _cell_border(cell, bottom=("single", 8, "081650"))
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        if ri == 0:
                            r.font.bold = True


# ------------------------------------------------------------ 封面
def _sect_break_par(doc, banner_section=True):
    """返回一个携带「封面节」sectPr 的空段落（零页边距，满幅题图）。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    ppr = p._p.get_or_add_pPr()
    sect = OxmlElement("w:sectPr")
    pg = OxmlElement("w:pgSz"); pg.set(qn("w:w"), "11906"); pg.set(qn("w:h"), "16838")
    mar = OxmlElement("w:pgMar")
    for k, v in (("top", 0), ("right", 0), ("bottom", 0), ("left", 0), ("header", 0), ("footer", 0), ("gutter", 0)):
        mar.set(qn(f"w:{k}"), str(v))
    sect.append(pg); sect.append(mar)
    ppr.append(sect)
    return p


def add_cover(doc, meta: dict, banner: Path):
    body = doc.element.body
    anchor = body[0]
    made = []

    def par(align=None, before=0, after=0, left=SIDE, right=SIDE, line=None):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.left_indent, pf.right_indent = left, right
        pf.first_line_indent = Pt(0)
        if line:
            pf.line_spacing = line
        if align is not None:
            p.alignment = align
        made.append(p._p)
        return p

    # 1 满幅题图
    p = par(left=0, right=0)
    p.add_run().add_picture(str(banner), width=Cm(21.0))
    # 2 单位行
    p = par(before=14, after=10)
    _run(p, "复旦大学", 13, True, BLUE, "黑体", spacing=40)
    _run(p, "　｜　", 13, False, GREY, "黑体")
    _run(p, meta.get("course", "功能学科综合实验"), 13, False, INK, "黑体", spacing=40)
    # 3 标题
    title = meta.get("title") or "实验报告"
    p = par(before=14, after=4, line=1.15)
    _run(p, title, 28, True, INK, "黑体", "Arial")
    # 4 英文副标
    p = par(after=2)
    _run(p, "LABORATORY REPORT", 11, True, TEAL, "黑体", "Arial", spacing=60)
    p = par(after=18)
    ppr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "18"); b.set(qn("w:space"), "1"); b.set(qn("w:color"), "149C8E")
    bd.append(b); ppr.append(bd)
    p.paragraph_format.right_indent = Cm(14.0)  # 短色条
    _run(p, "", 2)

    # 5 信息表
    rows = [("实验编号", meta.get("exp_no", "")), ("班　　级", meta.get("class", "")), ("姓　　名", meta.get("student", "")),
            ("学　　号", meta.get("student_id", "")), ("实验小组", meta.get("group", "")), ("实验日期", meta.get("date", "")),
            ("指导老师", meta.get("teacher", ""))]
    rows = [(k, v) for k, v in rows if k != "实验编号" or v]
    t = doc.add_table(rows=len(rows), cols=2)
    made.append(t._tbl)
    tblpr = t._tbl.tblPr
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), str(int(SIDE.cm * 567))); ind.set(qn("w:type"), "dxa"); tblpr.append(ind)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Cm(3.2), Cm(9.8)
    for ri, (k, v) in enumerate(rows):
        c0, c1 = t.rows[ri].cells
        c0.width, c1.width = Cm(3.2), Cm(9.8)
        for c in (c0, c1):
            _cell_border(c, bottom=("single", 4, "D5DBEE"))
            c.paragraphs[0].paragraph_format.space_before = Pt(5)
            c.paragraphs[0].paragraph_format.space_after = Pt(5)
        _run(c0.paragraphs[0], k, 11, False, GREY, "黑体")
        _run(c1.paragraphs[0], v if v else "＿＿＿＿＿＿＿＿", 13, bool(v), INK if v else GREY, "宋体")
    # 6 底部落款 + 分节符
    p = par(before=34, after=0)
    _run(p, "Fudan University · Integrated Functional Experiments", 9, False, GREY, "宋体", "Arial", spacing=20)
    last = _sect_break_par(doc)
    made.append(last._p)
    for el in made:
        anchor.addprevious(el)


# ------------------------------------------------------------ 页眉/页码
def fill_header(doc, text: str):
    for sec in doc.sections:
        if sec.header.is_linked_to_previous:
            continue
        for p in sec.header.paragraphs:
            for r in p.runs:
                if "{{HEADER}}" in r.text:
                    r.text = r.text.replace("{{HEADER}}", text)


def restart_numbering(doc):
    sect = doc.element.body.find(qn("w:sectPr"))
    if sect is not None and sect.find(qn("w:pgNumType")) is None:
        e = OxmlElement("w:pgNumType")
        e.set(qn("w:start"), "1")
        sect.append(e)
