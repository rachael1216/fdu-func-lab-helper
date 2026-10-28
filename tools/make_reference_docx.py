#!/usr/bin/env python3
"""make_reference_docx.py — 生成 Word 样式母版 templates/word/reference.docx

由 pandoc 默认 reference.docx 改造：设置页面（A4）、字体（宋体 + Times New Roman）、
标题/正文/题注/引用块/代码样式、页眉页脚（占位符在 build.py 中替换）。
一般同学不需要运行它；改版式时才改这里。  需要：pandoc、python-docx
"""
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "templates" / "word" / "reference.docx"
INK, BLUE, TEAL, GREY = RGBColor(0x08, 0x16, 0x50), RGBColor(0x1B, 0x3F, 0xB0), RGBColor(0x14, 0x9C, 0x8E), RGBColor(0x59, 0x60, 0x7A)
CN_BODY, CN_HEAD, CN_KAI, LATIN = "宋体", "黑体", "楷体", "Times New Roman"


def set_fonts(style_or_run, east=CN_BODY, latin=LATIN):
    el = style_or_run.font.element if hasattr(style_or_run, "font") and hasattr(style_or_run.font, "element") else None
    rpr = style_or_run._element.get_or_add_rPr() if hasattr(style_or_run, "_element") and hasattr(style_or_run._element, "get_or_add_rPr") else None
    if rpr is None:
        rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), latin)
    rfonts.set(qn("w:eastAsia"), east)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]


def para_border(style, side, color="149C8E", sz=8, space=2):
    ppr = style.element.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    bdr.append(b)


def get_style(doc, name, typ=WD_STYLE_TYPE.PARAGRAPH, base="Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, typ)
        if base:
            st.base_style = doc.styles[base]
        return st


def fmt(st, size, bold=False, color=None, east=CN_BODY, align=None, before=0, after=0, line=1.3, indent_first=None, italic=False, keep_next=False, latin=LATIN):
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.italic = italic
    if color is not None:
        st.font.color.rgb = color
    set_fonts(st, east, latin)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        pf.alignment = align
    if indent_first is not None:
        pf.first_line_indent = Pt(indent_first)


def add_field(par, instr):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)
    return r


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    base = OUT.with_suffix(".base.docx")
    with open(base, "wb") as fh:
        subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], stdout=fh, check=True)
    doc = Document(str(base))
    base.unlink()

    # ---- 页面
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin, s.bottom_margin = Cm(2.4), Cm(2.2)
    s.header_distance, s.footer_distance = Cm(1.2), Cm(1.1)

    # ---- 文本样式
    normal = doc.styles["Normal"]
    fmt(normal, 10.5, line=1.35, after=0)
    for n in ("Body Text", "First Paragraph"):
        st = get_style(doc, n)
        fmt(st, 10.5, line=1.4, after=3, indent_first=21)
    fmt(get_style(doc, "Compact"), 10.5, line=1.3, after=0, indent_first=0)
    fmt(get_style(doc, "Heading 1"), 14, True, INK, CN_HEAD, before=16, after=6, line=1.2, keep_next=True)
    fmt(get_style(doc, "Heading 2"), 12, True, BLUE, CN_HEAD, before=10, after=4, line=1.2, keep_next=True)
    fmt(get_style(doc, "Heading 3"), 11, True, RGBColor(0x2B, 0x33, 0x4D), CN_HEAD, before=8, after=3, line=1.2, keep_next=True)
    fmt(get_style(doc, "Heading 4"), 10.5, True, GREY, CN_HEAD, before=6, after=2, line=1.2, keep_next=True)
    para_border(doc.styles["Heading 1"], "bottom", "149C8E", 8, 3)
    fmt(get_style(doc, "Title"), 20, True, INK, CN_HEAD, WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4)
    fmt(get_style(doc, "Subtitle"), 12, False, GREY, CN_KAI, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    for n in ("Author", "Date", "Abstract"):
        fmt(get_style(doc, n), 10.5, False, GREY, CN_KAI, WD_ALIGN_PARAGRAPH.CENTER)

    # 自定义：报告标题块
    rt = get_style(doc, "ReportTitle", base="Normal")
    fmt(rt, 18, True, INK, CN_HEAD, WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4, line=1.25, indent_first=0)
    rm = get_style(doc, "ReportMeta", base="Normal")
    fmt(rm, 10.5, False, GREY, CN_KAI, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.4, indent_first=0)
    para_border(rm, "bottom", "D5DBEE", 6, 6)

    # 引用块（写作提示 / 重要说明）
    bq = get_style(doc, "Block Text")
    fmt(bq, 9.5, False, GREY, CN_KAI, before=3, after=3, line=1.3, indent_first=0)
    bq.paragraph_format.left_indent = Cm(0.5)
    bq.paragraph_format.right_indent = Cm(0.3)
    para_border(bq, "left", "149C8E", 18, 8)
    ppr = bq.element.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F7FA")
    ppr.append(shd)

    # 题注
    for n in ("Table Caption", "Image Caption", "Caption"):
        st = get_style(doc, n)
        fmt(st, 9.5, True, RGBColor(0x2B, 0x33, 0x4D), CN_KAI, WD_ALIGN_PARAGRAPH.CENTER, before=4, after=4, line=1.2, indent_first=0)
    doc.styles["Table Caption"].paragraph_format.keep_with_next = True
    for n in ("Figure", "Captioned Figure"):
        st = get_style(doc, n)
        fmt(st, 10.5, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=0, line=1.0, indent_first=0)
        st.paragraph_format.keep_with_next = True

    # 参考文献、代码
    fmt(get_style(doc, "Bibliography"), 9.5, line=1.25, after=2, indent_first=0)
    for n in ("Source Code", "Verbatim Char"):
        try:
            st = doc.styles[n]
        except KeyError:
            continue
        st.font.size = Pt(9)
        set_fonts(st, CN_BODY, "Consolas")
    try:
        doc.styles["Hyperlink"].font.color.rgb = BLUE
    except KeyError:
        pass
    # 表格文字
    tt = get_style(doc, "Table", WD_STYLE_TYPE.TABLE, base=None) if False else None  # 表格版式由 build.py 后处理

    # ---- 页眉页脚（占位符 {{HEADER}} 由 build.py 替换）
    hp = s.header.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    hp.text = ""
    run = hp.add_run("{{HEADER}}")
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    set_fonts(run, CN_KAI)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = hp._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "4"); b.set(qn("w:color"), "149C8E")
    bdr.append(b); ppr.append(bdr)

    fp = s.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for txt, field in (("— ", None), (None, "PAGE"), (" —", None)):
        if field:
            r = add_field(fp, field)
        else:
            r = fp.add_run(txt)
        r.font.size = Pt(9)
        r.font.color.rgb = GREY

    # 文档属性
    doc.core_properties.title = "功能学科综合实验报告"
    doc.core_properties.author = ""
    doc.save(str(OUT))
    print("✓", OUT)


if __name__ == "__main__":
    main()
