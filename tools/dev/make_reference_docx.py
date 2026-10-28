#!/usr/bin/env python3
"""make_reference_docx.py — 生成 Word 样式母版 templates/word/reference.docx

由 pandoc 默认 reference.docx 改造：设置页面（A4）、字体（宋体 + Times New Roman）、
标题/正文/题注/引用块/代码样式、页眉页脚（占位符在 build.py 中替换）。
一般同学不需要运行它；改版式时才改这里。  需要：pandoc、python-docx
"""
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent.parent
OUT = ROOT / "templates" / "word" / "reference.docx"
BLACK, SAGE, GREY = RGBColor(0, 0, 0), RGBColor(0x5F, 0x75, 0x68), RGBColor(0x59, 0x60, 0x6A)
CN_BODY, CN_HEAD, CN_KAI, LATIN = "宋体", "黑体", "楷体", "Times New Roman"


def set_fonts(style_or_run, east=CN_BODY, latin=LATIN):
    el = style_or_run.font.element if hasattr(style_or_run, "font") and hasattr(style_or_run.font, "element") else None
    rpr = style_or_run._element.get_or_add_rPr() if hasattr(style_or_run, "_element") and hasattr(style_or_run._element, "get_or_add_rPr") else None
    if rpr is None:
        rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), latin)
    rfonts.set(qn("w:eastAsia"), east)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]


def para_border(style, side, color="149C8E", sz=8, space=2):
    ppr = style.element.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    bdr.append(b)


def get_style(doc, name, typ=WD_STYLE_TYPE.PARAGRAPH, base="Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, typ)
        if base:
            st.base_style = doc.styles[base]
        return st


def fmt(st, size, bold=False, color=None, east=CN_BODY, align=None, before=0, after=0, line=1.3, indent_first=None, italic=False, keep_next=False, latin=LATIN):
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.italic = italic
    if color is not None:
        st.font.color.rgb = color
    set_fonts(st, east, latin)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        pf.alignment = align
    if indent_first is not None:
        pf.first_line_indent = Pt(indent_first)


def add_field(par, instr):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)
    return r


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    base = OUT.with_suffix(".base.docx")
    with open(base, "wb") as fh:
        subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], stdout=fh, check=True)
    doc = Document(str(base))
    base.unlink()

    # ---- 页面：A4，上下 2.54 cm，左右 3.0 cm（正文宽 15 cm）
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(3.0)
    s.top_margin, s.bottom_margin = Cm(2.54), Cm(2.54)
    s.header_distance, s.footer_distance = Cm(1.3), Cm(1.3)

    # ---- 文本样式（规范见 templates/FORMAT.md）
    fmt(doc.styles["Normal"], 12, line=1.5, after=0)
    for n in ("Body Text", "First Paragraph"):
        fmt(get_style(doc, n), 12, line=1.5, after=0, indent_first=24)
    fmt(get_style(doc, "Compact"), 12, line=1.4, after=0, indent_first=0)
    fmt(get_style(doc, "Heading 1"), 16, True, BLACK, CN_HEAD, before=18, after=8, line=1.2, keep_next=True)
    fmt(get_style(doc, "Heading 2"), 14, True, BLACK, CN_HEAD, before=12, after=6, line=1.2, keep_next=True)
    fmt(get_style(doc, "Heading 3"), 12, True, BLACK, CN_HEAD, before=8, after=4, line=1.2, keep_next=True)
    fmt(get_style(doc, "Heading 4"), 12, False, BLACK, CN_HEAD, before=6, after=2, line=1.2, keep_next=True)
    fmt(get_style(doc, "Title"), 22, True, BLACK, CN_HEAD, WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6)
    fmt(get_style(doc, "Subtitle"), 14, False, GREY, CN_KAI, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    for n in ("Author", "Date"):
        fmt(get_style(doc, n), 12, False, BLACK, CN_BODY, WD_ALIGN_PARAGRAPH.CENTER)
    fmt(get_style(doc, "Abstract"), 12, False, BLACK, CN_BODY, line=1.5, indent_first=24)
    # 报告标题块（正文首页）
    rt = get_style(doc, "ReportTitle", base="Normal")
    fmt(rt, 22, True, BLACK, CN_HEAD, WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, line=1.25, indent_first=0)
    rm = get_style(doc, "ReportMeta", base="Normal")
    fmt(rm, 10.5, False, GREY, CN_KAI, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.4, indent_first=0)
    # 目录
    for n in ("TOC Heading",):
        fmt(get_style(doc, n), 16, True, BLACK, CN_HEAD, WD_ALIGN_PARAGRAPH.CENTER, before=0, after=12, line=1.2, indent_first=0)
    # 引用块（写作提示 / AI 使用说明）
    bq = get_style(doc, "Block Text")
    fmt(bq, 10.5, False, GREY, CN_KAI, before=3, after=3, line=1.3, indent_first=0)
    bq.paragraph_format.left_indent = Cm(0.5)
    bq.paragraph_format.right_indent = Cm(0.3)
    para_border(bq, "left", "5F7568", 18, 8)
    ppr = bq.element.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F5F3")
    ppr.append(shd)
    # 图表题注：五号宋体居中；图题在图下、表题在表上
    for n in ("Table Caption", "Image Caption", "Caption"):
        st = get_style(doc, n)
        fmt(st, 10.5, False, BLACK, CN_BODY, WD_ALIGN_PARAGRAPH.CENTER, before=4, after=4, line=1.2, indent_first=0)
    doc.styles["Table Caption"].paragraph_format.keep_with_next = True
    for n in ("Figure", "Captioned Figure"):
        st = get_style(doc, n)
        fmt(st, 12, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=0, line=1.0, indent_first=0)
        st.paragraph_format.keep_with_next = True
    # 参考文献：五号，悬挂缩进
    bib = get_style(doc, "Bibliography")
    fmt(bib, 10.5, line=1.3, after=2, indent_first=0)
    bib.paragraph_format.left_indent = Pt(21)
    bib.paragraph_format.first_line_indent = Pt(-21)
    for n in ("Source Code", "Verbatim Char"):
        try:
            st = doc.styles[n]
        except KeyError:
            continue
        st.font.size = Pt(9.5)
        set_fonts(st, CN_BODY, "Consolas")
    try:
        doc.styles["Hyperlink"].font.color.rgb = BLACK
    except KeyError:
        pass

    # ---- 页眉：左 课程名，右 学号；底部细线。占位符由 build.py 替换
    hp = s.header.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = hp.paragraph_format
    pf.tab_stops.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT)
    pf.first_line_indent = Pt(0)
    for txt in ("{{HEADER_L}}", "\t", "{{HEADER_R}}"):
        run = hp.add_run(txt if txt != "\t" else "")
        if txt == "\t":
            run.add_tab()
        run.font.size = Pt(9)
        run.font.color.rgb = GREY
        set_fonts(run, CN_BODY)
    ppr = hp._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "4"); b.set(qn("w:color"), "5F7568")
    bdr.append(b); ppr.append(bdr)

    # ---- 页脚：居中页码（目录为罗马数字、正文为阿拉伯数字，由 build.py 分节设置）
    fp = s.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = add_field(fp, "PAGE")
    r.font.size = Pt(10.5)
    r.font.color.rgb = BLACK

    # 文档属性
    doc.core_properties.title = "功能学科综合实验报告"
    doc.core_properties.author = ""
    doc.save(str(OUT))
    print("✓", OUT)


if __name__ == "__main__":
    main()
